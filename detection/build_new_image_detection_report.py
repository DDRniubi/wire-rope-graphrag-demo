from __future__ import annotations

import argparse
import csv
import json
from datetime import datetime
from pathlib import Path
from typing import Any

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def load_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def load_descriptions(path: Path) -> dict[str, dict[str, Any]]:
    if not path.exists():
        return {}
    rows = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    return {str(row["image"]).replace("\\", "/"): row for row in rows}


def read_boxes(path: Path) -> list[dict[str, Any]]:
    boxes = []
    if not path.exists():
        return boxes
    for line in path.read_text(encoding="utf-8").splitlines():
        parts = line.split()
        if len(parts) >= 6:
            boxes.append({
                "class_id": int(parts[0]), "cx": float(parts[1]), "cy": float(parts[2]),
                "width": float(parts[3]), "height": float(parts[4]), "confidence": float(parts[5]),
            })
    return boxes


def merge(detection_dir: Path) -> list[dict[str, Any]]:
    descriptions = load_descriptions(detection_dir / "qwen_descriptions.jsonl")
    records = []
    for row in load_csv(detection_dir / "prediction_summary.csv"):
        rel = row["image"].replace("\\", "/")
        boxes = read_boxes(detection_dir / "labels" / Path(rel).with_suffix(".txt"))
        description = descriptions.get(rel, {})
        records.append({
            "image": rel,
            "status": row["status"],
            "yolo_detection_count": len(boxes),
            "yolo_max_confidence": max((box["confidence"] for box in boxes), default=0.0),
            "yolo_boxes": boxes,
            "result_image": str((detection_dir / "images" / rel).resolve()),
            "qwen_used": bool(description),
            "scene": description.get("scene", "未调用千问，采用YOLO离线检测"),
            "surface_condition": description.get("surface_condition", ""),
            "contamination": description.get("contamination", ""),
            "detection_difficulty": description.get("detection_difficulty", "未评估"),
            "short_caption": description.get("short_caption", f"YOLO检出断丝目标{len(boxes)}处。"),
            "detailed_caption": description.get("detailed_caption", ""),
            "manual_review_required": True,
        })
    return records


def write_data(detection_dir: Path, records: list[dict[str, Any]], args: argparse.Namespace) -> dict[str, Any]:
    with (detection_dir / "new_image_results.jsonl").open("w", encoding="utf-8") as handle:
        for row in records:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")
    fields = ["image", "status", "yolo_detection_count", "yolo_max_confidence", "qwen_used", "scene", "detection_difficulty", "short_caption", "result_image", "manual_review_required"]
    with (detection_dir / "new_image_results.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(records)
    summary = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "images_total": len(records),
        "images_with_detections": sum(row["yolo_detection_count"] > 0 for row in records),
        "detections_total": sum(row["yolo_detection_count"] for row in records),
        "average_max_confidence": round(sum(row["yolo_max_confidence"] for row in records) / max(1, len(records)), 4),
        "qwen_enabled": any(row["qwen_used"] for row in records),
        "model": str(Path(args.model).resolve()),
        "confidence_threshold": args.conf,
        "image_size": args.imgsz,
        "manual_review_required": True,
    }
    (detection_dir / "new_image_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def font(size: int):
    for path in (Path(r"C:\Windows\Fonts\simsun.ttc"), Path(r"C:\Windows\Fonts\simsun.ttf")):
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def contact_sheet(detection_dir: Path, records: list[dict[str, Any]]) -> Path:
    selected = sorted(records, key=lambda row: (row["yolo_detection_count"], row["yolo_max_confidence"]), reverse=True)[:6]
    cell_w, image_h, text_h = 720, 390, 86
    rows = max(1, (len(selected) + 1) // 2)
    sheet = Image.new("RGB", (cell_w * 2, (image_h + text_h) * rows), "white")
    draw = ImageDraw.Draw(sheet)
    for index, row in enumerate(selected):
        col, grid_row = index % 2, index // 2
        x, y = col * cell_w, grid_row * (image_h + text_h)
        image = Image.open(row["result_image"]).convert("RGB")
        image.thumbnail((cell_w - 16, image_h - 12), Image.Resampling.LANCZOS)
        sheet.paste(image, (x + (cell_w - image.width) // 2, y + (image_h - image.height) // 2))
        draw.rectangle((x, y, x + cell_w - 1, y + image_h + text_h - 1), outline="black", width=2)
        draw.text((x + 12, y + image_h + 8), Path(row["image"]).name[:38], fill="black", font=font(22))
        meta = f"候选框 {row['yolo_detection_count']}  最高置信度 {row['yolo_max_confidence']:.3f}"
        draw.text((x + 12, y + image_h + 44), meta, fill="black", font=font(20))
    path = detection_dir / "new_image_contact_sheet.jpg"
    sheet.save(path, quality=94)
    return path


def set_run(run, size=10.5, bold=False):
    run.font.name = "宋体"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), "宋体")


def paragraph(doc, text, size=10.5, bold=False, center=False, indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    if indent and not center:
        p.paragraph_format.first_line_indent = Cm(0.74)
    set_run(p.add_run(text), size, bold)
    return p


def heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run(text), 14, True)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, value in enumerate(headers):
        cell = t.rows[0].cells[col]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(value), 8.5, True)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "E7E6E6")
        cell._tc.get_or_add_tcPr().append(shd)
    for values in rows:
        cells = t.add_row().cells
        for col, value in enumerate(values):
            cells[col].text = ""
            cells[col].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[col].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col < 4 else WD_ALIGN_PARAGRAPH.LEFT
            set_run(p.add_run(str(value)), 8.2)
    for row in t.rows:
        for col, width in enumerate(widths):
            row.cells[col].width = Cm(width)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "BFBFBF")
        borders.append(node)
    t._tbl.tblPr.append(borders)


def report(detection_dir: Path, records: list[dict[str, Any]], summary: dict[str, Any], figure: Path) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin, section.bottom_margin = Cm(1.7), Cm(1.6)
    section.left_margin, section.right_margin = Cm(2.0), Cm(2.0)
    for name, size, bold in (("Normal", 10.5, False), ("Title", 18, True), ("Heading 1", 14, True)):
        style = doc.styles[name]
        style.font.name, style.font.size, style.font.bold = "宋体", Pt(size), bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{key}"), "宋体")
    title_border = doc.styles["Title"]._element.get_or_add_pPr().find(qn("w:pBdr"))
    if title_border is not None:
        doc.styles["Title"]._element.get_or_add_pPr().remove(title_border)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(title.add_run("钢丝绳新图片检测报告"), 18, True)
    paragraph(doc, f"生成时间：{summary['generated_at'].replace('T', ' ')}", 9.5, center=True, indent=False)
    heading(doc, "一 检测说明")
    mode = "YOLO检测加千问场景描述" if summary["qwen_enabled"] else "YOLO离线检测，未调用千问API"
    paragraph(doc, f"本次采用{mode}。断丝目标的位置、数量和置信度由YOLO模型输出。")
    heading(doc, "二 检测汇总")
    table(doc, ["图像数", "检出图像数", "候选框总数", "平均最高置信度"], [[summary["images_total"], summary["images_with_detections"], summary["detections_total"], f"{summary['average_max_confidence']:.3f}"]], [3.9] * 4)
    heading(doc, "三 代表性结果")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(figure), width=Cm(15.8))
    paragraph(doc, "图1  新图片YOLO检测结果", 9.2, center=True, indent=False)
    doc.add_page_break()
    heading(doc, "四 逐图结果摘要")
    rows = []
    for index, row in enumerate(sorted(records, key=lambda item: item["yolo_detection_count"], reverse=True)[:20], start=1):
        rows.append([index, Path(row["image"]).name, row["yolo_detection_count"], f"{row['yolo_max_confidence']:.3f}", row["short_caption"]])
    table(doc, ["序号", "图像", "框数", "最高置信度", "说明"], rows, [1.0, 4.0, 1.3, 2.1, 7.4])
    paragraph(doc, "完整框坐标和全部逐图结果保存在new_image_results.jsonl与new_image_results.csv中。", indent=False)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(footer.add_run("钢丝绳新图片检测报告"), 9)
    out = detection_dir / "钢丝绳新图片检测报告.docx"
    doc.save(out)
    return out


def main() -> None:
    parser = argparse.ArgumentParser(description="Build an offline/optional-Qwen report for newly supplied images.")
    parser.add_argument("--detection-dir", required=True)
    parser.add_argument("--model", required=True)
    parser.add_argument("--conf", type=float, default=0.25)
    parser.add_argument("--imgsz", type=int, default=1280)
    args = parser.parse_args()
    detection_dir = Path(args.detection_dir)
    records = merge(detection_dir)
    summary = write_data(detection_dir, records, args)
    figure = contact_sheet(detection_dir, records)
    output = report(detection_dir, records, summary, figure)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"report: {output.resolve()}")


if __name__ == "__main__":
    main()
